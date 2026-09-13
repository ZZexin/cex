"""Package artifact-authored N cells without round-tripping unrelated Excel XML."""
from pathlib import Path
import re
import zipfile
import xml.etree.ElementTree as ET
import openpyxl

source = Path(r'C:\Users\Zexin\Downloads\sno2zmo, 0.05ma,cm2dcg. DARK _27.0.0.1-BTS82-45-11-1-2818575313\sno2zmo\Selected\SNO2ZMO_10_G2.xlsx')
directory = Path(__file__).parent
authored = directory / 'energy_artifact_intermediate.xlsx'
output = directory / 'SNO2ZMO_10_G2_Energy修正版.xlsx'
ns = {'s': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
part = 'xl/worksheets/sheet1.xml'

with zipfile.ZipFile(authored) as generated, zipfile.ZipFile(source) as original:
    tree = ET.fromstring(generated.read(part))
    cells = {c.attrib['r']: c for c in tree.findall('.//s:sheetData/s:row/s:c', ns)}
    xml = original.read(part).decode('utf-8')
    replacements = 0
    def replace_cell(match):
        global replacements
        address = match.group(2)
        row = int(address[1:])
        if row < 3 or row == 748:
            return match.group(0)
        c = cells[address]
        formula = c.find('s:f', ns)
        value = c.find('s:v', ns)
        assert formula is not None and value is not None, address
        assert c.attrib.get('t', 'n') == 'n', address
        assert float(value.text) >= 0, address
        # Transfer only the new formula and cache. Keep original cell style/metadata.
        f = ET.Element('f')
        f.text = formula.text
        v = ET.Element('v')
        v.text = value.text
        replacements += 1
        return match.group(1) + ET.tostring(f, encoding='unicode') + ET.tostring(v, encoding='unicode') + '</c>'
    fixed = re.sub(r'(<c\b[^>]*\br="(N\d+)"[^>]*>).*?</c>', replace_cell, xml)
    assert replacements == 5496, replacements
    with zipfile.ZipFile(output, 'w', compression=zipfile.ZIP_DEFLATED) as target:
        for info in original.infolist():
            target.writestr(info, fixed.encode('utf-8') if info.filename == part else original.read(info.filename))

# Check all raw parts and cells outside the requested column are untouched.
with zipfile.ZipFile(source) as before, zipfile.ZipFile(output) as after:
    assert before.namelist() == after.namelist()
    for name in before.namelist():
        if name != part:
            assert before.read(name) == after.read(name), name
    strip_n = lambda data: re.sub(rb'<c\b[^>]*\br="N\d+"[^>]*>.*?</c>', b'', data)
    assert strip_n(before.read(part)) == strip_n(after.read(part))

raw = openpyxl.load_workbook(source, data_only=True).active
result = openpyxl.load_workbook(output, data_only=True).active
formulas = openpyxl.load_workbook(output, data_only=False).active
energy = 0.0
max_error = 0.0
for r in range(2, 5500):
    new_step = r == 2 or any(raw.cell(r,c).value != raw.cell(r-1,c).value for c in (2,4,5))
    if new_step:
        assert raw.cell(r,7).value == 0
        energy = 0.0
    else:
        energy += (raw.cell(r-1,8).value + raw.cell(r,8).value) / 2 * (raw.cell(r,13).value - raw.cell(r-1,13).value)
        expected = f'=IF(OR(B{r}<>B{r-1},D{r}<>D{r-1},E{r}<>E{r-1}),0,N{r-1}+(H{r-1}+H{r})/2*(M{r}-M{r-1}))'
        assert formulas.cell(r,14).value == expected
    actual = result.cell(r,14).value
    assert isinstance(actual, (int,float)), (r,actual)
    max_error = max(max_error, abs(actual-energy))
    assert abs(actual-energy) < 1e-10, (r,actual,energy)
print('Verified all 5498 data rows; all unrelated XML and cells unchanged; max error', max_error)
print('N747 charge endpoint:', result['N747'].value, 'mWh')
print('N5499 discharge endpoint:', result['N5499'].value, 'mWh')
print(ascii(str(output)))
